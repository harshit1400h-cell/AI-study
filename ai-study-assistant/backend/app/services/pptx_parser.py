"""PowerPoint (.pptx) parser using python-pptx."""

import os
from typing import Any, Dict, List
from pptx import Presentation


def parse_pptx(file_path: str) -> List[Dict[str, Any]]:
    """Extract text from PowerPoint presentations slide-by-slide.

    Args:
        file_path: Absolute or relative path to the PPTX document.

    Returns:
        List of dictionaries with 'text' and 'metadata'.

    Raises:
        FileNotFoundError: If the file does not exist.
        ValueError: If the presentation cannot be opened.
    """
    if not os.path.isfile(file_path):
        raise FileNotFoundError(f"PPTX file not found: {file_path}")

    filename = os.path.basename(file_path)
    try:
        prs = Presentation(file_path)
    except Exception as e:
        raise ValueError(f"Failed to open PPTX file {filename}: {e}") from e

    results: List[Dict[str, Any]] = []

    for idx, slide in enumerate(prs.slides):
        slide_number = idx + 1
        slide_title = ""

        # Attempt to extract title from slide title shape
        try:
            if slide.shapes.title and slide.shapes.title.has_text_frame:
                slide_title = slide.shapes.title.text.strip()
        except Exception:
            slide_title = ""

        slide_texts: List[str] = []

        for shape in slide.shapes:
            # Avoid repeating the slide title twice if already captured
            if shape == getattr(slide.shapes, "title", None):
                continue

            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    line = para.text.strip()
                    if line:
                        slide_texts.append(line)
            elif shape.has_table:
                for row in shape.table.rows:
                    row_text = [c.text.strip() for c in row.cells if c.text.strip()]
                    if row_text:
                        slide_texts.append(" | ".join(row_text))

        combined_body = "\n".join(slide_texts).strip()
        full_text = f"Slide Title: {slide_title}\n{combined_body}" if slide_title else combined_body

        if not full_text.strip():
            continue

        section_name = slide_title if slide_title else f"Slide {slide_number}"

        results.append({
            "text": full_text.strip(),
            "metadata": {
                "filename": filename,
                "page": slide_number,
                "slide_number": slide_number,
                "slide_title": slide_title,
                "section": section_name,
            },
        })

    return results

