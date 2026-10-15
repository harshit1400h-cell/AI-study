"""DOCX parser using python-docx with heading detection and section grouping."""

import os
from typing import Any, Dict, List
import docx


def parse_docx(file_path: str) -> List[Dict[str, Any]]:
    """Extract and group text from a Word (.docx) document by headings/sections.

    Args:
        file_path: Absolute or relative path to the DOCX document.

    Returns:
        List of dictionaries with 'text' and 'metadata'.

    Raises:
        FileNotFoundError: If the file does not exist.
        ValueError: If the file cannot be parsed.
    """
    if not os.path.isfile(file_path):
        raise FileNotFoundError(f"DOCX file not found: {file_path}")

    filename = os.path.basename(file_path)
    try:
        doc = docx.Document(file_path)
    except Exception as e:
        raise ValueError(f"Failed to open DOCX file {filename}: {e}") from e

    sections: List[Dict[str, Any]] = []
    current_heading = "General"
    current_texts: List[str] = []
    current_start_index = 0

    def flush_section(heading: str, texts: List[str], start_idx: int) -> None:
        combined = "\n\n".join(t.strip() for t in texts if t.strip())
        if combined:
            sections.append({
                "text": combined,
                "metadata": {
                    "filename": filename,
                    "section": heading,
                    "paragraph_index": start_idx,
                    "page": 1,
                },
            })

    for idx, para in enumerate(doc.paragraphs):
        text = para.text.strip()
        if not text:
            continue

        style_name = para.style.name.lower() if para.style and para.style.name else ""
        is_heading = "heading" in style_name or "title" in style_name

        if is_heading:
            # Flush previous section
            if current_texts:
                flush_section(current_heading, current_texts, current_start_index)
                current_texts = []
            current_heading = text
            current_start_index = idx
        else:
            current_texts.append(text)

    # Also extract any text from tables in the document
    if doc.tables:
        table_texts: List[str] = []
        for table in doc.tables:
            for row in table.rows:
                row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_cells:
                    table_texts.append(" | ".join(row_cells))
        if table_texts:
            current_texts.append("Tables:\n" + "\n".join(table_texts))

    # Flush final accumulated section
    if current_texts:
        flush_section(current_heading, current_texts, current_start_index)

    return sections

