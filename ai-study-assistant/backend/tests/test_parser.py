"""Unit tests for document parsers (PDF, DOCX, PPTX, TXT)."""

import os
import tempfile
import pytest
import docx
import fitz
from pptx import Presentation
from pptx.util import Inches

from app.services.docx_parser import parse_docx
from app.services.pdf_parser import parse_pdf
from app.services.pptx_parser import parse_pptx
from app.services.txt_parser import parse_txt


def test_txt_parser():
    """Test plain text parser with multiple sections."""
    sample_text = (
        "Introduction to Photosynthesis\n"
        "Photosynthesis is the process used by plants to convert light energy into chemical energy.\n\n"
        "Light-Dependent Reactions\n"
        "These reactions take place in the thylakoid membranes of chloroplasts."
    )
    with tempfile.NamedTemporaryFile("w", suffix=".txt", delete=False, encoding="utf-8") as f:
        f.write(sample_text)
        temp_path = f.name

    try:
        results = parse_txt(temp_path)
        assert len(results) == 2
        assert "Photosynthesis" in results[0]["text"]
        assert results[0]["metadata"]["page"] == 1
        assert "Light-Dependent" in results[1]["text"]
    finally:
        if os.path.exists(temp_path):
            os.remove(temp_path)


def test_pdf_parser():
    """Test PDF parser with PyMuPDF generated document."""
    doc = fitz.open()
    page1 = doc.new_page()
    page1.insert_text((50, 50), "Chapter 1: Quantum Mechanics Basics.\nWave-particle duality is fundamental.")
    page2 = doc.new_page()
    page2.insert_text((50, 50), "Chapter 2: Schrödinger Equation.\nThe time-dependent wave equation.")

    with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as f:
        temp_path = f.name

    try:
        doc.save(temp_path)
        doc.close()

        results = parse_pdf(temp_path)
        assert len(results) == 2
        assert "Quantum Mechanics" in results[0]["text"]
        assert results[0]["metadata"]["page"] == 1
        assert results[0]["metadata"]["total_pages"] == 2
        assert "Schrödinger" in results[1]["text"]
        assert results[1]["metadata"]["page"] == 2
    finally:
        if os.path.exists(temp_path):
            os.remove(temp_path)


def test_docx_parser():
    """Test DOCX parser with python-docx generated document."""
    doc = docx.Document()
    doc.add_heading("Cell Biology", level=1)
    doc.add_paragraph("The cell is the basic structural and functional unit of all living organisms.")
    doc.add_heading("Mitochondria", level=2)
    doc.add_paragraph("Mitochondria generate most of the chemical energy needed to power biochemical reactions.")

    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as f:
        temp_path = f.name

    try:
        doc.save(temp_path)
        results = parse_docx(temp_path)
        assert len(results) >= 2
        assert any("cell is the basic" in r["text"] for r in results)
        assert any("Mitochondria generate" in r["text"] for r in results)
    finally:
        if os.path.exists(temp_path):
            os.remove(temp_path)


def test_pptx_parser():
    """Test PPTX parser with python-pptx generated presentation."""
    prs = Presentation()
    blank_layout = prs.slide_layouts[6]
    slide1 = prs.slides.add_slide(blank_layout)

    txBox = slide1.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(2))
    tf = txBox.text_frame
    tf.text = "Newton's Laws of Motion"
    p2 = tf.add_paragraph()
    p2.text = "First Law: Inertia states an object at rest stays at rest."

    with tempfile.NamedTemporaryFile(suffix=".pptx", delete=False) as f:
        temp_path = f.name

    try:
        prs.save(temp_path)
        results = parse_pptx(temp_path)
        assert len(results) == 1
        assert "Newton's Laws" in results[0]["text"]
        assert results[0]["metadata"]["slide_number"] == 1
    finally:
        if os.path.exists(temp_path):
            os.remove(temp_path)


def test_file_not_found():
    """Verify FileNotFoundError is raised when file does not exist."""
    with pytest.raises(FileNotFoundError):
        parse_txt("non_existent_file_12345.txt")

    with pytest.raises(FileNotFoundError):
        parse_pdf("non_existent_file_12345.pdf")

    with pytest.raises(FileNotFoundError):
        parse_docx("non_existent_file_12345.docx")

    with pytest.raises(FileNotFoundError):
        parse_pptx("non_existent_file_12345.pptx")

